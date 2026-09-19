"""Field extraction from SI / BL attachments.

Formats: .txt (label lines), .pdf (coordinate de-interleave + regex for the
bottom single-string lines), .xlsx (label/value columns), .docx (2-col table
with bilingual Chinese glosses).

Label synonyms mirror the generator's pools.LABELS: the same field is drawn
with different labels on SI vs BL ("Load Port" == "Port of Loading").
"""
from __future__ import annotations

import os
import re

# Canonical field -> label variants, longest-first matching downstream.
SYNONYMS: dict[str, list[str]] = {
    "shipper": ["Shipper/Exporter", "Shipper (Principal or Seller)", "Shipper", "SHIPPER"],
    "consignee": ["Consignee (Non-Negotiable)", "Consignee", "CONSIGNEE", "To the Order of"],
    "notify_party": ["Notify Party/Intermediate Consignee", "Notify Party", "NOTIFY PARTY", "Notify"],
    "port_of_loading": ["Port of Loading (POL)", "PORT OF LOADING", "Port of Loading", "Load Port", "POL"],
    "port_of_discharge": ["Port of Discharge (POD)", "PORT OF DISCHARGE", "Port of Discharge", "Discharge Port", "POD"],
    "container_count": ["No. of Containers or Packages", "No. of Containers", "Total Containers", "Container Count"],
    "gross_weight_kg": ["Gross Weight (KG)", "Gross Wt (kgs)", "GROSS WEIGHT", "Gross Weight毛重(KGS)"],
}

_FLAT_LABELS = sorted(
    ((lab.upper(), field) for field, labs in SYNONYMS.items() for lab in labs),
    key=lambda t: -len(t[0]),
)

COMPARE_FIELDS = ["shipper", "consignee", "notify_party", "port_of_loading",
                  "port_of_discharge", "container_count", "gross_weight_kg"]

NAME_FIELDS = {"shipper", "consignee", "notify_party"}
_COMPANY_SUFFIX_RE = re.compile(
    r"\b(LTD|LIMITED|INC|CO\.?|CORP|GMBH|LLC|PLC)\b\.?", re.I)
_ADDRESS_KEYWORD_RE = re.compile(
    r"\b(ROAD|STREET|AVENUE|FLOOR|NO\.|UNIT|BLOCK|SUITE)\b", re.I)


def _split_name_from_address(value: str) -> str:
    """collapse_lines can merge a company name with the address line that
    follows it. Cut right after a recognized company suffix; if none,
    cut right before a recognized address keyword."""
    m = _COMPANY_SUFFIX_RE.search(value)
    if m:
        return value[:m.end()].strip().rstrip(",")
    m = _ADDRESS_KEYWORD_RE.search(value)
    if m:
        return value[:m.start()].strip().rstrip(",")
    return value

EXTRACTABLE_EXTS = {".txt", ".pdf", ".xlsx", ".docx"}


def _norm_text(v: str) -> str:
    return re.sub(r"\s+", " ", v).strip()


# ---------------------------------------------------------------------------
# Text-line parsing (shared by txt / xlsx / docx renderers)
# ---------------------------------------------------------------------------
def parse_fields(text: str) -> dict[str, str]:
    """Parse 'Label: value' lines into {canonical_field: raw_value}."""
    fields: dict[str, str] = {}
    for raw in text.splitlines():
        line = raw.strip()
        if not line or set(line) <= {"="}:
            continue
        up = line.upper()
        # PDF total row: 'TOTAL Gross Weight (KG): ...' -- only strip 'TOTAL '
        # when a gross-weight label follows (never before 'Total Containers'!)
        if up.startswith("TOTAL ") and any(
            up[6:].startswith(lab) for lab, f in _FLAT_LABELS if f == "gross_weight_kg"
        ):
            up, line = up[6:], line[6:]
        field, rest = _match_label(up)
        if not field:
            continue
        field, rest = _match_label(up)
        if not field:
            continue
        val = line[len(line) - len(rest):].lstrip(": ").strip()
        if field in NAME_FIELDS:
            val = _split_name_from_address(val)
        fields[field] = val
    return fields


def _match_label(line_upper: str):
    for lab, field in _FLAT_LABELS:
        if line_upper.startswith(lab):
            rest = line_upper[len(lab):]
            if rest[:1] in (":", " ", ""):
                return field, rest
    return None, None


# ---------------------------------------------------------------------------
# Per-format extractors -> pseudo 'Label: value' text
# ---------------------------------------------------------------------------
def extract_txt(path: str) -> str:
    with open(path, encoding="utf-8") as f:
        return f.read()


def extract_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    lines = []
    for row in ws.iter_rows(values_only=True):
        label = str(row[0]) if row and row[0] is not None else ""
        value = str(row[1]) if row and len(row) > 1 and row[1] is not None else ""
        value = value.split("|")[0]  # xlsx packs 'NAME | addr; addr' in one cell
        if label:
            lines.append(f"{label}: {value}")
    wb.close()
    return "\n".join(lines)


def extract_docx(path: str) -> str:
    import docx
    d = docx.Document(path)
    lines = []
    for table in d.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            label = row.cells[0].text.strip()
            value = row.cells[1].text.strip().splitlines()[0] if row.cells[1].text.strip() else ""
            # docx labels carry a Chinese gloss: 'Shipper (发货人)'
            label = re.sub(r"\s*\([^()]*\)\s*$", "", label)
            if label:
                lines.append(f"{label}: {value}")
    return "\n".join(lines)


def extract_pdf(path: str) -> str | None:
    """Pseudo 'Label: value' lines for a generated PDF.

    Block labels are drawn at x=20mm and values at x=60mm; a long label
    ('Notify Party/Intermediate Consignee') runs INTO the value zone, so raw
    text extraction interleaves both strings. We de-interleave on character
    coordinates. Gross weight / container count only exist as bottom
    single-string lines ('TOTAL Gross Weightnn(KGS): 23,702 KG' -- the CJK in
    one label variant renders as garbage glyphs), handled by regex.
    Returns None when the page has no text layer (image-only scan).
    """
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        if not pdf.pages:
            raise ValueError("no pages")
        page = pdf.pages[0]
        text = page.extract_text() or ""
        if not text.strip() or not page.chars:
            return None

        lines = []
        m = re.search(r"(?im)^(?:total\s+)?gross\s*(?:weight|wt).*?:\s*([\d,]+(?:\.\d+)?)\s*kgs?\s*$", text)
        if m:
            lines.append(f"Gross Weight (KG): {m.group(1)}")
        m = re.search(
            r"(?im)^(?:no\.?\s*of\s*containers(?:\s*or\s*packages)?|total\s*containers|container\s*count)\s*:\s*(\d+)\s*x",
            text,
        )
        if m:
            lines.append(f"Container Count: {m.group(1)} x")

        VALUE_X = 170.1  # 60mm: where values start
        TOL = 3.5
        buckets: dict[float, list] = {}
        for c in page.chars:
            buckets.setdefault(round(c["top"], 1), []).append(c)
        for cs in buckets.values():
            cs.sort(key=lambda c: c["x0"])
            zone = "".join(c["text"] for c in cs if c["x0"] < VALUE_X - 2)
            zn = re.sub(r"[^A-Z0-9]", "", zone.upper())
            if not zn:
                continue
            cand = None
            for lab, field in _FLAT_LABELS:
                if field in ("container_count", "gross_weight_kg"):
                    continue  # bottom lines only, handled by regex above
                ln = re.sub(r"[^A-Z0-9]", "", lab.upper())
                if ln and (zn.startswith(ln) or ln.startswith(zn)):
                    if cand is None or len(ln) > len(cand[0]):
                        cand = (ln, field)
            if not cand:
                continue
            ln, field = cand
            expected = cs[0]["x0"]  # label run starts at 20mm
            got = 0
            value = []
            for c in cs:
                ch = c["text"]
                if got < len(ln) and abs(c["x0"] - expected) < TOL:
                    if ch.isalnum():
                        if ch.upper() != ln[got]:
                            value.append(ch)  # value char sitting in label zone
                            continue
                        got += 1
                    expected = c["x0"] + c["width"]
                else:
                    value.append(ch)
            val = "".join(value).strip().lstrip("):.,;")  # drop stray label ')'
            if val:
                lines.append(f"{SYNONYMS[field][0]}: {val}")
        return "\n".join(lines)


_EXTRACTORS = {".txt": extract_txt, ".pdf": extract_pdf,
               ".xlsx": extract_xlsx, ".docx": extract_docx}


# ---------------------------------------------------------------------------
# Document loading
# ---------------------------------------------------------------------------
class DocumentStatus:
    TEXT = "text"
    UNREADABLE = "unreadable"


def load_document(data_root: str, rel_path: str):
    """Return (status, text). status: 'text' or 'unreadable'.

    'unreadable' covers: missing/empty file, corrupt bytes, and PDFs without
    a text layer (scans) -- the pipeline must escalate these, not guess.
    """
    full = os.path.join(data_root, rel_path)
    ext = os.path.splitext(rel_path)[1].lower()
    try:
        if not os.path.exists(full) or os.path.getsize(full) == 0:
            return DocumentStatus.UNREADABLE, ""
        if ext not in EXTRACTABLE_EXTS:
            return DocumentStatus.UNREADABLE, ""
        text = _EXTRACTORS[ext](full)
        if text is None:
            return DocumentStatus.UNREADABLE, ""
        return DocumentStatus.TEXT, text
    except Exception:
        return DocumentStatus.UNREADABLE, ""


def doc_title(text: str) -> str:
    """First heading line (used for wrong-doc-type detection on .txt)."""
    for line in text.splitlines():
        line = line.strip()
        if not line or set(line) <= {"="}:
            continue
        return line.upper()
    return ""
